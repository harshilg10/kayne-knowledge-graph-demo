"""Raw file contract round trip: manifest + bodies -> raw files (render_docs) -> DocumentRecords (ingest).

Everything renders into temporary directories from a temporary copy of the authored bodies; a placeholder
stands in for any body not written yet. authoring/document_bodies and data/unstructured are never written.
"""
from __future__ import annotations

import ast
import csv
import email
import email.policy
import json
import re
import shutil
from collections import Counter
from dataclasses import dataclass
from email.utils import format_datetime
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from kgdemo import config, ingest, render_docs
from kgdemo import world as w
from kgdemo.catalog import normalize_text
from kgdemo.manifest import DocSpec, doc_specs
from kgdemo.models import DocumentRecord

HERO_TEXT = (
    "Investment Committee Meeting\n\nDate: 2026-01-15\n\nParticipants:\nDouglas Whitfield\nMeredith Calloway\n\n"
    "Discussion:\nWe remain bullish on Nvidia and other AI infrastructure providers.\n"
    "The technology team believes demand for GPU capacity will continue through 2027.\n"
)
SPECS = {s.doc_id: s for s in doc_specs()}
TITLES = {doc_id: s.title for doc_id, s in SPECS.items()}
PEOPLE = {p["name"]: p for p in w.people()}
METADATA_KEYS = {
    "Meeting Notes": {"participants", "meeting_type"},
    "IC Memo": {"authors", "decision_ref", "for_meeting_date", "cited_titles"},
    "Research Report": {"authors"},
    "Due Diligence": {"authors", "borrower"},
    "Email": {"sender", "sender_email", "to", "cc", "to_emails", "cc_emails"},
    "CRM Note": {"client_id", "client_name", "owner", "crm_type"},
}


@dataclass
class Corpus:
    raw_dir: Path
    counts: dict[str, int]
    records: dict[str, DocumentRecord]
    bodies: dict[str, str]          # the body each document was rendered from


def collapse(text: str) -> str:
    return " ".join(text.split())


def placeholder_body(spec: DocSpec) -> str:
    """Stand-in for a body not authored yet: the brief plus a sentence listing the required mentions."""
    text = spec.brief or f"Placeholder body for {spec.doc_id}."
    if spec.must_mention:
        text += f"\n\nThis document mentions {', '.join(spec.must_mention)}."
    return text.encode("ascii", "replace").decode("ascii") + "\n"


def raw_path(raw_dir: Path, spec: DocSpec) -> Path:
    return raw_dir / render_docs.FOLDER_OF[spec.doc_type] / render_docs.file_name(spec)


def expected_org(name: str) -> str | None:
    person = PEOPLE[name]
    return None if person["person_type"] == "employee" else person["organization"]


@pytest.fixture(scope="module")
def corpus(tmp_path_factory) -> Corpus:
    source_dir = tmp_path_factory.mktemp("source_docs")
    for path in sorted(config.SOURCE_DOCS_DIR.glob("*.txt")):
        shutil.copyfile(path, source_dir / path.name)
    for spec in SPECS.values():
        target = source_dir / f"{spec.doc_id}.txt"
        if spec.fixed_body is None and (not target.exists() or not target.read_text(encoding="utf-8-sig").strip()):
            target.write_text(placeholder_body(spec), encoding="utf-8")
    raw_dir = tmp_path_factory.mktemp("raw")
    counts = render_docs.render_all(raw_dir, source_dir, require_all=True)
    records = {r.doc_id: r for r in ingest.ingest_all(raw_dir)}
    bodies = {doc_id: spec.fixed_body if spec.fixed_body is not None
              else (source_dir / f"{doc_id}.txt").read_text(encoding="utf-8-sig") for doc_id, spec in SPECS.items()}
    return Corpus(raw_dir, counts, records, bodies)


# ---- rendering ---------------------------------------------------------------------
def test_every_spec_is_rendered_and_ingested(corpus):
    expected = Counter(render_docs.FOLDER_OF[s.doc_type] for s in SPECS.values())
    assert corpus.counts == {folder: expected[folder] for folder in ingest.FOLDERS}
    assert sum(corpus.counts.values()) == len(SPECS) == 101
    assert sorted(corpus.records) == sorted(SPECS)


def test_file_names(corpus):
    expected = {render_docs.FOLDER_OF[s.doc_type]: set() for s in SPECS.values()}
    for spec in SPECS.values():
        if spec.doc_type != "CRM Note":
            expected[render_docs.FOLDER_OF[spec.doc_type]].add(render_docs.file_name(spec))
    expected["crm"] = {ingest.CRM_EXPORT_NAME}
    for folder, names in expected.items():
        assert {p.name for p in (corpus.raw_dir / folder).iterdir()} == names
    for spec in SPECS.values():
        slug = render_docs.slugify(spec.title)
        assert 0 < len(slug) <= 60 and slug == slug.strip("-") and "--" not in slug
        assert set(slug) <= set("abcdefghijklmnopqrstuvwxyz0123456789-")


def test_slugify():
    assert render_docs.slugify("GPU demand outlook - follow-ups from yesterday's IC") == \
        "gpu-demand-outlook-follow-ups-from-yesterday-s-ic"
    assert render_docs.slugify("Power & Energy Research Meeting") == "power-energy-research-meeting"
    assert render_docs.slugify("Private Credit Fund II - Capital Call Notice #6") == \
        "private-credit-fund-ii-capital-call-notice-6"
    long_slug = render_docs.slugify("Reducing overlap between Global Technology Fund and AI Infrastructure Fund")
    assert long_slug == "reducing-overlap-between-global-technology-fund-and-ai-infra"
    assert render_docs.slugify("A" * 59 + " B") == "a" * 59           # cut at 60, trailing '-' stripped


def test_render_all_replaces_only_its_folders_and_checks_bodies_first(tmp_path):
    raw = tmp_path / "raw"
    stale = raw / "meeting_notes" / "DOC-OLD-01__stale.md"
    stale.parent.mkdir(parents=True)
    stale.write_text("stale")
    (raw / "other").mkdir()
    (raw / "other" / "keep.txt").write_text("keep")
    (raw / "readme.txt").write_text("keep")
    empty_sources = tmp_path / "source_docs"
    empty_sources.mkdir()
    with pytest.raises(FileNotFoundError, match="DOC-MN-01"):
        render_docs.render_all(raw, empty_sources, require_all=True)
    assert stale.exists()                                   # nothing is deleted when bodies are missing
    counts = render_docs.render_all(raw, empty_sources, require_all=False)
    assert counts == {**dict.fromkeys(ingest.FOLDERS, 0), "meeting_notes": 1}    # the hero has a fixed body
    assert not stale.exists()
    assert [p.name for p in (raw / "meeting_notes").iterdir()] == ["DOC-MN-03__investment-committee-meeting.md"]
    assert (raw / "other" / "keep.txt").read_text() == "keep" and (raw / "readme.txt").read_text() == "keep"
    assert [r.doc_id for r in ingest.ingest_all(raw)] == ["DOC-MN-03"]   # the empty CRM export has no rows


# ---- the round trip: every spec ----------------------------------------------------------
def expected_body(spec: DocSpec, body: str) -> str:
    """The body as rendered: an IC memo also lists its supporting documents at the end."""
    if spec.doc_type == "IC Memo" and spec.cites:
        body += "\n\n" + "\n".join([ingest.SUPPORTING_DOCUMENTS, *(TITLES[c] for c in spec.cites)])
    return normalize_text(body)


def check_header_facts(spec: DocSpec, md: dict) -> None:
    if spec.doc_type == "Meeting Notes":
        meeting = w.meeting_by_key(spec.meeting_key)
        assert [p["name"] for p in md["participants"]] == meeting["participants"]
        assert [p["organization"] for p in md["participants"]] == [expected_org(n) for n in meeting["participants"]]
        assert md["meeting_type"] == meeting["meeting_type"]
    elif spec.doc_type == "IC Memo":
        assert md["authors"] == spec.authors
        assert md["decision_ref"] == spec.decision_ref
        assert md["for_meeting_date"] == spec.for_meeting_date
        assert md["cited_titles"] == [TITLES[c] for c in spec.cites]
    elif spec.doc_type == "Research Report":
        assert md["authors"] == spec.authors
    elif spec.doc_type == "Due Diligence":
        assert md["authors"] == spec.authors
        assert md["borrower"] == spec.borrower
    elif spec.doc_type == "Email":
        assert (md["sender"], md["to"], md["cc"]) == (spec.sender, spec.to, spec.cc)
        assert md["sender_email"] == PEOPLE[spec.sender]["email"]
        assert md["to_emails"] == [PEOPLE[n]["email"] for n in spec.to]
        assert md["cc_emails"] == [PEOPLE[n]["email"] for n in spec.cc]
    else:
        assert (md["client_id"], md["owner"], md["crm_type"]) == (spec.client_id, spec.authors[0], spec.crm_type)
        assert md["client_name"] == w.client_by_id(spec.client_id)["name"]
    if spec.doc_type in ("IC Memo", "Research Report", "Due Diligence"):
        orgs = {a: expected_org(a) for a in spec.authors if expected_org(a)}
        assert md.get("author_organizations", {}) == orgs


@pytest.mark.parametrize("doc_id", sorted(SPECS))
def test_round_trip(corpus, doc_id):
    spec, record = SPECS[doc_id], corpus.records[doc_id]
    assert (record.doc_id, record.doc_type, record.date, record.title, record.source_system) == \
        (spec.doc_id, spec.doc_type, spec.date, spec.title, spec.source_system)
    assert collapse(record.body) == collapse(expected_body(spec, corpus.bodies[doc_id]))
    check_header_facts(spec, record.metadata)
    optional = {"author_organizations"} if spec.doc_type in ("IC Memo", "Research Report", "Due Diligence") else set()
    assert METADATA_KEYS[spec.doc_type] <= set(record.metadata) <= METADATA_KEYS[spec.doc_type] | optional
    # canonical text: title line, header lines, a blank line, the body; no letterhead or disclaimer
    assert record.text.startswith(spec.title + "\n")
    assert record.text[:record.body_offset].endswith("\n\n") and record.text.endswith("\n")
    assert ingest.LETTERHEAD not in record.text and ingest.DISCLAIMER not in record.text
    if spec.doc_type == "CRM Note":
        assert record.source_path.endswith(f"crm/{ingest.CRM_EXPORT_NAME}#{doc_id}")
    else:
        assert record.source_path.endswith("/" + render_docs.file_name(spec))


def test_pdf_and_docx_bodies_keep_line_structure(corpus):
    """Every source line comes back as exactly one line: wrapped PDF lines are rejoined, paragraphs never merge."""
    for spec in SPECS.values():
        if spec.doc_type not in ("IC Memo", "Research Report", "Due Diligence"):
            continue
        record = corpus.records[spec.doc_id]
        source = expected_body(spec, corpus.bodies[spec.doc_id])
        assert [line for line in record.body.splitlines() if line] == \
            [collapse(line) for line in source.splitlines() if line.strip()], spec.doc_id
        assert "\n\n\n" not in record.text


def test_text_formats_keep_the_body_verbatim(corpus):
    """Meeting notes, emails and CRM notes keep the authored line structure exactly."""
    for spec in SPECS.values():
        if spec.doc_type in ("Meeting Notes", "Email", "CRM Note"):
            body = normalize_text(corpus.bodies[spec.doc_id])
            expected = "\n".join(line.rstrip() for line in body.splitlines()).strip() + "\n"
            assert corpus.records[spec.doc_id].body == expected, spec.doc_id


def test_meeting_note_text_is_the_file_content(corpus):
    for spec in SPECS.values():
        if spec.doc_type == "Meeting Notes":
            content = raw_path(corpus.raw_dir, spec).read_text()
            assert corpus.records[spec.doc_id].text == content
            assert ("\nMeeting Type: " in content) == (spec.doc_id != config.HERO_DOC_ID)


# ---- the hero note -------------------------------------------------------------------------
def test_hero_note(corpus):
    spec = SPECS[config.HERO_DOC_ID]
    path = raw_path(corpus.raw_dir, spec)
    assert path.name == "DOC-MN-03__investment-committee-meeting.md"
    assert path.read_text() == HERO_TEXT
    record = corpus.records[config.HERO_DOC_ID]
    assert record.text == HERO_TEXT
    assert record.body_offset == HERO_TEXT.index("Discussion:")
    assert record.body.startswith("Discussion:\nWe remain bullish on Nvidia")
    assert record.metadata["participants"] == [{"name": "Douglas Whitfield", "organization": None},
                                               {"name": "Meredith Calloway", "organization": None}]
    assert record.metadata["meeting_type"] == "Investment Committee"
    assert (record.title, record.date, record.source_system) == \
        ("Investment Committee Meeting", "2026-01-15", "Document Management")


# ---- parse_meeting_notes (also used for notes pasted into the app) ------------------------------
def test_parse_meeting_notes_accepts_one_line_participants():
    note = ("Investment Committee Meeting\nDate: 2026-01-15\nParticipants: Douglas Whitfield, Meredith Calloway\n\n"
            "Discussion:\nWe remain bullish on Nvidia.\n")
    record = ingest.parse_meeting_notes(note, "DOC-LIVE-TEST", "pasted")
    assert [p["name"] for p in record.metadata["participants"]] == ["Douglas Whitfield", "Meredith Calloway"]
    assert record.metadata["meeting_type"] == "Investment Committee"
    assert (record.title, record.date, record.source_path) == ("Investment Committee Meeting", "2026-01-15", "pasted")
    assert record.text == note and record.body == "Discussion:\nWe remain bullish on Nvidia.\n"


def test_parse_meeting_notes_tolerates_pasted_variants():
    note = ("Hartwell Check-in\r\nDate: September 20, 2026\r\nParticipants: Michael Torres; "
            "Robert Hartwell (Hartwell Family Trust) and Leah Goldberg (Pemberton Grid Consulting, LLC)\r\n\r\n"
            "Robert\u2019s question \u2013 Nvidia volatility.   \r\n")
    record = ingest.parse_meeting_notes(note, "DOC-LIVE-TEST", "pasted")
    assert record.metadata == {
        "participants": [{"name": "Michael Torres", "organization": None},
                         {"name": "Robert Hartwell", "organization": "Hartwell Family Trust"},
                         {"name": "Leah Goldberg", "organization": "Pemberton Grid Consulting, LLC"}],
        "meeting_type": "General"}
    assert record.date == "2026-09-20"
    assert record.body == "Robert's question - Nvidia volatility.\n"     # normalized, LF, no trailing spaces


def test_parse_meeting_notes_participant_list_ends_at_a_label_or_sentence():
    for first_body_line in ("Discussion:", "We remain bullish on Nvidia."):
        note = ("Ad Hoc Meeting\nDate: 2026-01-15\nParticipants:\nDouglas Whitfield\nMeredith Calloway\n"
                f"{first_body_line}\nMore text.\n")
        record = ingest.parse_meeting_notes(note, "DOC-LIVE-TEST", "pasted")
        assert [p["name"] for p in record.metadata["participants"]] == ["Douglas Whitfield", "Meredith Calloway"]
        assert record.body == f"{first_body_line}\nMore text.\n"


def test_parse_meeting_notes_rejects_notes_without_title_or_date():
    with pytest.raises(ValueError, match="Date"):
        ingest.parse_meeting_notes("Some Meeting\n\nDiscussion:\nText.\n", "DOC-LIVE-TEST", "pasted")
    with pytest.raises(ValueError, match="title"):
        ingest.parse_meeting_notes("Date: 2026-01-15\n\nDiscussion:\nText.\n", "DOC-LIVE-TEST", "pasted")
    with pytest.raises(ValueError, match="empty"):
        ingest.parse_meeting_notes(" \n\n", "DOC-LIVE-TEST", "pasted")


# ---- format details ------------------------------------------------------------------------
def test_ic_memo_docx_layout(corpus):
    spec = SPECS["DOC-ICM-03"]
    doc = Document(str(raw_path(corpus.raw_dir, spec)))
    paragraphs = [p.text for p in doc.paragraphs]
    assert paragraphs[:8] == [
        "ARBOR PEAK WEALTH PARTNERS", "INVESTMENT COMMITTEE MEMORANDUM", "To: Investment Committee",
        "From: Meredith Calloway, Brendan Kessler", "Date: 2026-02-05", f"Subject: {spec.title}", "Decision Reference: IC-2026-004",
        "For Meeting: Investment Committee Meeting, 2026-02-12"]
    assert all(run.bold for p in doc.paragraphs[:2] for run in p.runs)
    cited = [TITLES[c] for c in spec.cites]
    assert paragraphs[-len(cited) - 2:] == [ingest.SUPPORTING_DOCUMENTS, *cited, ingest.DISCLAIMER]
    assert "" not in paragraphs
    uncited = Document(str(raw_path(corpus.raw_dir, SPECS["DOC-ICM-01"])))
    assert ingest.SUPPORTING_DOCUMENTS not in [p.text for p in uncited.paragraphs]
    assert corpus.records["DOC-ICM-01"].metadata["cited_titles"] == []


def test_pdf_layout(corpus):
    dd = PdfReader(raw_path(corpus.raw_dir, SPECS["DOC-DD-02"]))
    lines = "\n".join(page.extract_text() for page in dd.pages).splitlines()
    assert lines[0] == ingest.DD_BANNER
    assert "Prepared by: Leah Goldberg (Pemberton Grid Consulting)" in lines
    assert lines[-1] == ingest.DISCLAIMER
    assert corpus.records["DOC-DD-02"].metadata["author_organizations"] == {"Leah Goldberg": "Pemberton Grid Consulting"}
    research = PdfReader(raw_path(corpus.raw_dir, SPECS["DOC-RR-01"]))
    first = research.pages[0].extract_text().splitlines()
    assert first[0] == ingest.RESEARCH_BANNER
    assert "Authors: Meredith Calloway, Arjun Mehta" in first and "Published: 2025-11-20" in first


def test_wrapped_pdf_title_and_paragraphs_are_rebuilt(tmp_path):
    title = ("A Deliberately Long Research Title About GPU Capacity, Hyperscaler Capex and Power Demand "
             "Through 2027")
    paragraph = " ".join(["Demand for GPU capacity keeps running ahead of supply - and the gap persists."] * 6)
    body = f"Summary:\n{paragraph}\n\nKey Points:\n- First point.\n- Second point - with a dash.\n\nClosing line.\n"
    spec = DocSpec("DOC-RR-99", "Research Report", "2026-05-01", title, authors=["Meredith Calloway", "Leah Goldberg"])
    path = tmp_path / "research" / render_docs.file_name(spec)
    path.parent.mkdir()
    render_docs.write_research_pdf(spec, body, path)
    pdf_lines = PdfReader(path).pages[0].extract_text().splitlines()
    assert title not in pdf_lines and paragraph not in pdf_lines            # both really wrap in the PDF
    [record] = ingest.ingest_path(path)
    assert record.title == title
    assert record.text == (f"{title}\nAuthors: Meredith Calloway, Leah Goldberg (Pemberton Grid Consulting)\n"
                           f"Published: 2026-05-01\n\n{body}")
    assert record.body == body
    assert record.metadata == {"authors": ["Meredith Calloway", "Leah Goldberg"],
                               "author_organizations": {"Leah Goldberg": "Pemberton Grid Consulting"}}


def test_email_headers(corpus):
    msg = email.message_from_bytes(raw_path(corpus.raw_dir, SPECS["DOC-EM-01"]).read_bytes(),
                                   policy=email.policy.default)
    assert msg["From"] == "Michael Torres <michael.torres@arborpeakwealth.com>"
    assert msg["To"] == "Robert Hartwell <robert.hartwell@hartwellfamily.com>"
    assert msg["Cc"] == "Elaine Hartwell <elaine.hartwell@hartwellfamily.com>"
    assert msg["Date"] == format_datetime(render_docs.sent_at(SPECS["DOC-EM-01"]))
    assert msg["Date"].endswith("-0500")                     # Dallas, daylight saving time in August
    assert msg["Message-ID"] == render_docs.message_id(SPECS["DOC-EM-01"])
    assert msg["Subject"] == SPECS["DOC-EM-01"].title
    message_ids = set()
    for spec in SPECS.values():
        if spec.doc_type != "Email":
            continue
        msg = email.message_from_bytes(raw_path(corpus.raw_dir, spec).read_bytes(), policy=email.policy.default)
        sent = msg["Date"].datetime
        assert sent.date().isoformat() == spec.date and sent.utcoffset().total_seconds() in (-5 * 3600, -6 * 3600,
                                                                                            -7 * 3600, -8 * 3600)
        assert 7 <= sent.hour <= 18
        assert re.fullmatch(r"<[A-Z]{2}\dPR\d{2}MB\d{4}[0-9A-F]{32}@[A-Z]{2}\dPR\d{2}MB\d{4}\.namprd\d{2}\."
                            r"prod\.outlook\.com>", msg["Message-ID"])
        message_ids.add(msg["Message-ID"])
        assert ("Cc" in msg) == bool(spec.cc)
    assert len(message_ids) == sum(s.doc_type == "Email" for s in SPECS.values())


def test_crm_export(corpus):
    with (corpus.raw_dir / "crm" / ingest.CRM_EXPORT_NAME).open(newline="") as fh:
        reader = csv.DictReader(fh)
        rows = list(reader)
    assert tuple(reader.fieldnames) == ingest.CRM_COLUMNS
    crm_ids = sorted(doc_id for doc_id, s in SPECS.items() if s.doc_type == "CRM Note")
    assert [row["NoteId"] for row in rows] == crm_ids
    for row in rows:
        spec = SPECS[row["NoteId"]]
        assert (row["CreatedDate"], row["ClientId"], row["Owner"], row["Type"], row["Subject"]) == \
            (spec.date, spec.client_id, spec.authors[0], spec.crm_type, spec.title)


# ---- contract and CLI ----------------------------------------------------------------------
def test_duplicate_doc_ids_are_rejected(corpus, tmp_path):
    folder = tmp_path / "raw" / "meeting_notes"
    folder.mkdir(parents=True)
    hero = raw_path(corpus.raw_dir, SPECS[config.HERO_DOC_ID])
    shutil.copyfile(hero, folder / "DOC-MN-03__first.md")
    shutil.copyfile(hero, folder / "DOC-MN-03__second.md")
    with pytest.raises(ValueError, match="DOC-MN-03"):
        ingest.ingest_all(tmp_path / "raw")


def test_source_paths_are_relative_to_the_project_root():
    assert ingest._relative(config.UNSTRUCTURED_DIR / "emails" / "DOC-EM-01__x.eml") == "data/unstructured/emails/DOC-EM-01__x.eml"


def test_ingest_never_reads_the_manifest_or_world():
    tree = ast.parse(Path(ingest.__file__).read_text())
    modules = {node.module or "" for node in ast.walk(tree) if isinstance(node, ast.ImportFrom)}
    modules |= {alias.name for node in ast.walk(tree) if isinstance(node, (ast.Import, ast.ImportFrom))
                for alias in node.names}
    assert not any(m.split(".")[-1] in ("manifest", "world") for m in modules)


def test_ingest_cli_writes_documents_json(corpus, tmp_path, capsys):
    out = tmp_path / "documents.json"
    assert ingest.main(["--raw-dir", str(corpus.raw_dir), "--out", str(out)]) == 0
    records = [DocumentRecord.from_dict(d) for d in json.loads(out.read_text())]
    assert [r.doc_id for r in records] == sorted(SPECS)
    assert records[0] == corpus.records[records[0].doc_id]
    printed = capsys.readouterr().out
    assert "Meeting Notes" in printed and "101" in printed
