"""Renders the raw "as found" document files from the manifest specs and the authored bodies.

    python -m kgdemo.render_docs                   # authoring/document_bodies -> data/unstructured
    python -m kgdemo.render_docs --allow-missing   # skip documents whose body is not written yet

Header facts (dates, people, subjects, references) come from the manifest and the world model; only the
body is authored. Every format follows the raw file contract in docs/ARCHITECTURE.md, whose constants live
in ingest.py - the parser that reads these files back without access to the manifest.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import re
import shutil
from datetime import date, datetime, timedelta
from email.headerregistry import Address
from email.message import EmailMessage
from email.utils import format_datetime
from functools import lru_cache
from pathlib import Path
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from docx import Document
from docx.document import Document as WordDocument
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate

from . import config
from . import world as w
from .ingest import (CRM_COLUMNS, CRM_EXPORT_NAME, DD_BANNER, DISCLAIMER, FOLDERS, LETTERHEAD, MEMO_HEADING,
                     RESEARCH_BANNER, SUPPORTING_DOCUMENTS)
from .manifest import DocSpec, doc_specs

FOLDER_OF = {doc_type: folder for folder, (doc_type, _) in FOLDERS.items()}
EXTENSION_OF = {doc_type: ext for doc_type, ext in FOLDERS.values()}
HOME_TZ = "America/Chicago"                     # head office (Dallas); advisors send from their own office
OFFICE_TZ = {"Dallas": "America/Chicago", "Houston": "America/Chicago", "Chicago": "America/Chicago",
             "San Antonio": "America/Chicago", "Los Angeles": "America/Los_Angeles", "Seattle": "America/Los_Angeles"}
SLUG_MAX = 60


# =============================================================================
# public API
# =============================================================================
def render_all(out_dir: Path = config.UNSTRUCTURED_DIR, source_dir: Path = config.SOURCE_DOCS_DIR,
               require_all: bool = True) -> dict[str, int]:
    """Renders every document into out_dir/<folder>/ and returns the number of documents per folder.

    The six raw folders are deleted and recreated, so no stale file survives; nothing else in out_dir is
    touched. Bodies are loaded before anything is deleted: a missing (or blank) body raises
    FileNotFoundError when require_all, otherwise that document is skipped. CRM notes are counted as
    rows of the single export file.
    """
    specs = doc_specs()
    bodies = {s.doc_id: load_body(s, source_dir) for s in specs}
    missing = [s.doc_id for s in specs if bodies[s.doc_id] is None]
    if missing and require_all:
        raise FileNotFoundError(f"{len(missing)} document bodies are missing or empty in {source_dir}: "
                                f"{', '.join(missing)}")
    out_dir = Path(out_dir)
    for folder in FOLDERS:
        if (out_dir / folder).exists():
            shutil.rmtree(out_dir / folder)
        (out_dir / folder).mkdir(parents=True)
    titles = {s.doc_id: s.title for s in specs}
    counts = dict.fromkeys(FOLDERS, 0)
    crm_rows: list[dict[str, str]] = []
    for spec in specs:
        body = bodies[spec.doc_id]
        if body is None:
            continue
        counts[FOLDER_OF[spec.doc_type]] += 1
        if spec.doc_type == "CRM Note":
            crm_rows.append(crm_row(spec, body))
        else:
            render_document(spec, body, out_dir / FOLDER_OF[spec.doc_type] / file_name(spec), titles)
    write_crm_export(crm_rows, out_dir / FOLDER_OF["CRM Note"] / CRM_EXPORT_NAME)
    return counts


def render_document(spec: DocSpec, body: str, path: Path, titles: dict[str, str]) -> None:
    """Writes one non-CRM document; `titles` maps doc ids to titles for an IC memo's supporting documents."""
    if spec.doc_type == "Meeting Notes":
        path.write_text(meeting_notes_text(spec, body), encoding="utf-8")
    elif spec.doc_type == "IC Memo":
        write_ic_memo(spec, body, [titles[doc_id] for doc_id in spec.cites], path)
    elif spec.doc_type == "Research Report":
        write_research_pdf(spec, body, path)
    elif spec.doc_type == "Due Diligence":
        write_due_diligence_pdf(spec, body, path)
    elif spec.doc_type == "Email":
        path.write_bytes(email_bytes(spec, body))
    else:
        raise ValueError(f"{spec.doc_id}: no renderer for {spec.doc_type}")


def load_body(spec: DocSpec, source_dir: Path = config.SOURCE_DOCS_DIR) -> str | None:
    """The spec's fixed body, else the authored body file; None when that file is missing or blank."""
    if spec.fixed_body is not None:
        return spec.fixed_body
    path = Path(source_dir) / f"{spec.doc_id}.txt"
    if not path.is_file():
        return None
    body = path.read_text(encoding="utf-8-sig")
    return body if body.strip() else None


def slugify(title: str) -> str:
    """Lowercase title, runs of non-alphanumerics -> '-', at most SLUG_MAX characters."""
    return re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")[:SLUG_MAX].rstrip("-")


def file_name(spec: DocSpec) -> str:
    return f"{spec.doc_id}__{slugify(spec.title)}{EXTENSION_OF[spec.doc_type]}"


# =============================================================================
# people
# =============================================================================
@lru_cache(maxsize=1)
def _people() -> dict[str, dict]:
    return {p["name"]: p for p in w.people()}


def display_name(name: str) -> str:
    """Employees appear by name; client contacts and external people as "Name (Organization)"."""
    person = _people()[name]
    return name if person["person_type"] == "employee" else f"{name} ({person['organization']})"


def _names(names: list[str]) -> str:
    return ", ".join(display_name(n) for n in names)


# =============================================================================
# meeting notes (.md)
# =============================================================================
def meeting_notes_text(spec: DocSpec, body: str) -> str:
    if spec.meeting_key is None:
        raise ValueError(f"{spec.doc_id}: meeting notes need a meeting_key")
    meeting = w.meeting_by_key(spec.meeting_key)
    lines = [spec.title, "", f"Date: {spec.date}"]
    if spec.doc_id != config.HERO_DOC_ID:      # the brief's note has no type line; ingestion infers it
        lines.append(f"Meeting Type: {meeting['meeting_type']}")
    lines += ["", "Participants:", *(display_name(n) for n in meeting["participants"]), "", _clean_body(body)]
    return "\n".join(lines) + "\n"


# =============================================================================
# IC memo (.docx)
# =============================================================================
_BRAND_RGB = RGBColor(0x1F, 0x3A, 0x5F)
_MUTED_RGB = RGBColor(0x59, 0x59, 0x59)


def write_ic_memo(spec: DocSpec, body: str, cited_titles: list[str], path: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    _docx_paragraph(doc, LETTERHEAD, bold=True, size=14, color=_BRAND_RGB, space_after=2)
    _docx_paragraph(doc, MEMO_HEADING, bold=True, color=_BRAND_RGB, space_after=12)
    fields = [("To", "Investment Committee"), ("From", _names(spec.authors)), ("Date", spec.date),
              ("Subject", spec.title), ("Decision Reference", spec.decision_ref),
              ("For Meeting", spec.for_meeting_date and f"Investment Committee Meeting, {spec.for_meeting_date}")]
    fields = [(key, value) for key, value in fields if value]
    for i, (key, value) in enumerate(fields):
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"{key}:").bold = True
        paragraph.add_run(f" {value}")
        paragraph.paragraph_format.space_after = Pt(14 if i == len(fields) - 1 else 0)
    for line in _body_lines(body):
        label = _is_label(line)
        _docx_paragraph(doc, line, bold=label, space_before=8 if label else 0)
    if cited_titles:
        _docx_paragraph(doc, SUPPORTING_DOCUMENTS, bold=True, space_before=10, space_after=2)
        for title in cited_titles:
            _docx_paragraph(doc, title, italic=True, space_after=2)
    _docx_paragraph(doc, DISCLAIMER, italic=True, size=8, color=_MUTED_RGB, space_before=18)
    props = doc.core_properties
    props.title, props.subject, props.comments = spec.title, spec.decision_ref or "", ""
    props.author = props.last_modified_by = ", ".join(spec.authors)
    h = _digest(spec)
    props.created = datetime.fromisoformat(spec.date).replace(hour=8 + h[9] % 9, minute=h[10] % 60)
    props.modified = props.created + timedelta(minutes=12 + h[11] % 90)
    doc.save(str(path))


def _docx_paragraph(doc: WordDocument, text: str, *, bold: bool = False, italic: bool = False,
                    size: float | None = None, color: RGBColor | None = None, space_before: float = 0,
                    space_after: float = 6) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)


# =============================================================================
# research and due diligence (.pdf)
# =============================================================================
_BRAND = colors.HexColor("#1F3A5F")
_MUTED = colors.HexColor("#595959")
# Words never break across lines (no hyphenation, no long-word splitting), so text extraction can rejoin
# wrapped lines with a single space. Paragraph spacing keeps the baselines of consecutive paragraphs at
# least 1.5x the font size apart while wrapped lines sit closer: that is how ingest finds paragraphs.
_UNBROKEN = {"hyphenationLang": "", "embeddedHyphenation": 0, "splitLongWords": 0}
_PDF_STYLES = {
    "banner": ParagraphStyle("banner", fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=_BRAND,
                             **_UNBROKEN),
    "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=18, leading=22, spaceBefore=14,
                            spaceAfter=10, **_UNBROKEN),
    "meta": ParagraphStyle("meta", fontName="Helvetica", fontSize=10, leading=13, spaceAfter=5, textColor=_MUTED,
                           **_UNBROKEN),
    "body": ParagraphStyle("body", fontName="Helvetica", fontSize=10.5, leading=14, spaceAfter=6, **_UNBROKEN),
    "label": ParagraphStyle("label", fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=8,
                            spaceAfter=6, **_UNBROKEN),
    "disclaimer": ParagraphStyle("disclaimer", fontName="Helvetica-Oblique", fontSize=8, leading=10,
                                 textColor=_MUTED, **_UNBROKEN),
}


def write_research_pdf(spec: DocSpec, body: str, path: Path) -> None:
    meta = [f"Authors: {_names(spec.authors)}", f"Published: {spec.date}"]
    _write_pdf(path, RESEARCH_BANNER, spec.title, meta, body, author=", ".join(spec.authors))


def write_due_diligence_pdf(spec: DocSpec, body: str, path: Path) -> None:
    meta = [f"Borrower: {spec.borrower}", f"Prepared by: {_names(spec.authors)}", f"Date: {spec.date}"]
    _write_pdf(path, DD_BANNER, spec.title, meta, body, author=", ".join(spec.authors))


def _write_pdf(path: Path, banner: str, title: str, meta: list[str], body: str, author: str) -> None:
    doc = SimpleDocTemplate(str(path), pagesize=LETTER, leftMargin=inch, rightMargin=inch, topMargin=inch,
                            bottomMargin=inch, title=title, author=author, creator=config.FIRM_NAME, invariant=1)
    story = [Paragraph(escape(banner), _PDF_STYLES["banner"]),
             HRFlowable(width="100%", thickness=0.8, color=_BRAND, spaceBefore=3, spaceAfter=0),
             Paragraph(escape(title), _PDF_STYLES["title"]),
             *(Paragraph(escape(line), _PDF_STYLES["meta"]) for line in meta),
             HRFlowable(width="100%", thickness=0.4, color=_MUTED, spaceBefore=4, spaceAfter=12),
             *_pdf_body(body),
             HRFlowable(width="100%", thickness=0.4, color=_MUTED, spaceBefore=14, spaceAfter=4),
             Paragraph(escape(DISCLAIMER), _PDF_STYLES["disclaimer"])]
    doc.build(story)


def _pdf_body(body: str) -> list:
    """One paragraph per body line. Each paragraph (with any labels above it) is kept on one page."""
    flow, group = [], []
    for line in _body_lines(body):
        label = _is_label(line)
        group.append(Paragraph(escape(line), _PDF_STYLES["label" if label else "body"]))
        if not label:
            flow.append(KeepTogether(group))
            group = []
    if group:
        flow.append(KeepTogether(group))
    return flow


# =============================================================================
# email (.eml)
# =============================================================================
def email_bytes(spec: DocSpec, body: str) -> bytes:
    msg = EmailMessage()
    msg["From"] = _address(spec.sender)
    msg["To"] = [_address(n) for n in spec.to]
    if spec.cc:
        msg["Cc"] = [_address(n) for n in spec.cc]
    msg["Date"] = format_datetime(sent_at(spec))
    msg["Subject"] = spec.title
    msg["Message-ID"] = message_id(spec)
    msg.set_content(_clean_body(body) + "\n")
    return bytes(msg)


def _digest(spec: DocSpec) -> bytes:
    return hashlib.sha256(f"{spec.doc_id}|{spec.date}|{spec.title}".encode()).digest()


def sent_at(spec: DocSpec) -> datetime:
    """Deterministic send time on the document date, between 07:30 and 18:29 in the sender's office time zone."""
    advisor = next((a for a in w.ADVISORS if a["id"] == _people()[spec.sender]["id"]), None) if spec.sender else None
    zone = ZoneInfo(OFFICE_TZ.get(advisor["office"], HOME_TZ) if advisor else HOME_TZ)
    h = _digest(spec)
    minute_of_day = 7 * 60 + 30 + int.from_bytes(h[:2], "big") % (11 * 60)
    day = date.fromisoformat(spec.date)
    return datetime(day.year, day.month, day.day, minute_of_day // 60, minute_of_day % 60, h[2] % 60, tzinfo=zone)


def message_id(spec: DocSpec) -> str:
    """Exchange Online style Message-ID, deterministic per document."""
    h = _digest(spec)
    server = f"{('SA', 'DM', 'BY', 'CH')[h[3] % 4]}{h[4] % 9 + 1}PR{h[5] % 19 + 1:02d}MB{4000 + int.from_bytes(h[6:8], 'big') % 6000}"
    return f"<{server}{h[8:24].hex().upper()}@{server}.namprd{h[5] % 19 + 1:02d}.prod.outlook.com>"


def _address(name: str | None) -> Address:
    if name is None:
        raise ValueError("email participant missing")
    return Address(display_name=name, addr_spec=_people()[name]["email"])


# =============================================================================
# CRM export (.csv)
# =============================================================================
def crm_row(spec: DocSpec, body: str) -> dict[str, str]:
    return {"NoteId": spec.doc_id, "CreatedDate": spec.date, "ClientId": spec.client_id,
            "ClientName": w.client_by_id(spec.client_id)["name"], "Owner": spec.authors[0], "Type": spec.crm_type,
            "Subject": spec.title, "Body": _clean_body(body)}


def write_crm_export(rows: list[dict[str, str]], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=CRM_COLUMNS)
        writer.writeheader()
        writer.writerows(sorted(rows, key=lambda row: row["NoteId"]))


# =============================================================================
# body text helpers
# =============================================================================
def _clean_body(body: str) -> str:
    """LF line endings, no trailing spaces, no leading/trailing blank lines."""
    return "\n".join(line.rstrip() for line in body.replace("\r\n", "\n").split("\n")).strip()


def _body_lines(body: str) -> list[str]:
    """Non-empty body lines: each becomes one paragraph in .docx and .pdf files."""
    return [line.strip() for line in body.splitlines() if line.strip()]


def _is_label(line: str) -> bool:
    """A section label such as "Investment Thesis:" (rendered bold)."""
    return line.endswith(":") and len(line) <= 60 and not line.startswith("- ")


# =============================================================================
# CLI
# =============================================================================
def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Render the raw document files from the manifest and bodies.")
    parser.add_argument("--out-dir", type=Path, default=config.UNSTRUCTURED_DIR)
    parser.add_argument("--source-dir", type=Path, default=config.SOURCE_DOCS_DIR)
    parser.add_argument("--allow-missing", action="store_true",
                        help="skip documents whose body file is missing instead of failing")
    args = parser.parse_args(argv)
    counts = render_all(args.out_dir, args.source_dir, require_all=not args.allow_missing)
    for folder, n in counts.items():
        print(f"{folder:<14} {n:>4}")
    total = sum(counts.values())
    print(f"{'total':<14} {total:>4}  -> {args.out_dir}")
    skipped = len(doc_specs()) - total
    if skipped:
        print(f"skipped {skipped} documents without a body")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
